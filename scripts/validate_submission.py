"""Validate the generated submission files.

    .venv-docs/Scripts/python scripts/validate_submission.py

Checks team details, title, faculty, group, leftover template instructions, references, images,
slide count, text overflow heuristics, PDF page count and that the original templates are unchanged.
Exit code 0 = all checks passed (unresolved administrative placeholders are reported, not failed).
"""

from __future__ import annotations

import hashlib
import re
import sys
from pathlib import Path

import fitz
from docx import Document
from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parents[1]
SUB = ROOT / "submission"
DOCX = SUB / "Distributed_CDN_Simulator_Final_Report.docx"
PDF = SUB / "Distributed_CDN_Simulator_Final_Report.pdf"
PPTX = SUB / "Distributed_CDN_Simulator_Final_Presentation.pptx"
TEMPLATES = {  # SHA-256 recorded before any generation
    ROOT / "Final_Computer_Networks_Mini_Project_Report_Template.docx":
        "20922173d091fc34e88bd11090d38f8bb025bc53ca616506a2ddf02fe6ce848a",
    ROOT / "MiniProject_PPT_Template_TY_CSE_Final.pptx":
        "1fe0bdbdf179313c209812459b796ee4bad5504beb88af9e1fc09cc30a86910b",
}
TITLE = "Distributed CDN Simulator with Latency-Aware Routing and Adaptive Caching"
TEAM = [("More Geet Dattatraya", "1262250006"), ("Shinde Soham Manoj", "1262250009"),
        ("Jagtap Prasad Sachin", "1262250018"), ("Gaikwad Shreyash Sachin", "1262250005")]
INSTRUCTION_PATTERNS = [r"\[PROJECT TITLE\]", r"\[Student \d", r"\[PRN\]", r"\[Faculty Name\]", r"\[FACULTY NAME\]",
                        r"\[GROUP NUMBER\]", r"\[Write ", r"\[Insert ", r"\[INSERT", r"Using This Template",
                        r"Formatting and Submission", r"\[Objective\]", r"\[Test / input\]", r"\[Tool / device\]",
                        r"\[Measurable", r"\[What networking", r"\[Name", r"\[Keyword", r"\[Author", r"\[Relevant",
                        r"Replace all bracketed", r"\[NAME\]", r"\[No\.\]"]
ALLOWED_PLACEHOLDERS = re.compile(r"\[(PANEL|BATCH|Panel|Batch|Date|DD MONTH YYYY) – to be confirmed\]|"
                                  r"\[To be completed by the team\]|\[File / test / section\]|"
                                  r"\[Name – responsibility, to be completed by the team\]|"
                                  r"\[Contributions are not recorded in the repository – to be completed\]")

results: list[tuple[str, bool, str]] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    results.append((name, ok, detail))


def docx_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    for s in doc.sections:
        parts.extend(p.text for p in s.footer.paragraphs)
    return "\n".join(parts)


def validate_docx() -> None:
    doc = Document(str(DOCX))
    text = docx_text(doc)
    check("DOCX: project title present", text.count(TITLE) >= 2, f"{text.count(TITLE)} occurrences")
    check("DOCX: faculty Dr. Swati", "Dr. Swati" in text)
    check("DOCX: Group No. 2", "Group No. 2" in text)
    check("DOCX: academic year 2026–2027", "2026–2027" in text)
    for name, prn in TEAM:
        n = sum(1 for t in doc.tables for row in t.rows
                if any(name in c.text for c in row.cells) and any(prn in c.text for c in row.cells))
        check(f"DOCX: {name} with PRN {prn} in team tables", n >= 4, f"{n} table rows (cover, certificate, declaration, contributions)")
    leftovers = [p for p in INSTRUCTION_PATTERNS if re.search(p, ALLOWED_PLACEHOLDERS.sub("", text))]
    check("DOCX: no template instructions left", not leftovers, ", ".join(leftovers))
    unresolved = sorted(set(ALLOWED_PLACEHOLDERS.findall(text)) | set(m.group(0) for m in ALLOWED_PLACEHOLDERS.finditer(text)))
    check("DOCX: only expected administrative placeholders remain", True,
          "; ".join(sorted(set(m.group(0) for m in ALLOWED_PLACEHOLDERS.finditer(text)))))
    empty_headings = [p for p in doc.paragraphs if p.style.name.startswith("Heading") and not p.text.strip()]
    check("DOCX: no empty headings", not empty_headings, f"{len(empty_headings)} found")
    h = [p.text for p in doc.paragraphs if p.style.name in ("Heading 1",)]
    expected = ["Acknowledgement", "Abstract", "Table of Contents", "List of Figures", "List of Tables",
                "1 Introduction and Project Objectives", "2 Proposed Design and Requirements",
                "3 Implementation and Technical Working", "4 Testing and Results", "5 Discussion and Conclusion",
                "References", "AI Usage Declaration", "Appendices"]
    check("DOCX: chapter structure", h == expected, " | ".join(h))
    refs = [p.text for p in doc.paragraphs if re.match(r"^\[\d+\]\t", p.text)]
    body = "\n".join(p.text for p in doc.paragraphs if not re.match(r"^\[\d+\]\t", p.text))
    cited = {int(n) for grp in re.findall(r"\[(\d+(?:\],\s*\[\d+)*)\]", body) for n in re.findall(r"\d+", grp)}
    check("DOCX: every reference is cited", set(range(1, len(refs) + 1)) <= cited,
          f"{len(refs)} references; uncited: {sorted(set(range(1, len(refs) + 1)) - cited)}")
    check("DOCX: no citation without a reference", cited <= set(range(1, len(refs) + 1)), f"cited {sorted(cited)}")
    captions = "\n".join(p.text for p in doc.paragraphs if p.style.name in ("Figure Caption", "Table Caption"))
    figs = re.findall(r"^Figure ([\dA-Z]+\.\d+):", captions, re.M)
    tabs = re.findall(r"^Table ([\dA-Z]+\.\d+):", captions, re.M)
    check("DOCX: figure captions unique", len(figs) == len(set(figs)), f"{len(figs)} figures")
    check("DOCX: table captions unique", len(tabs) == len(set(tabs)), f"{len(tabs)} tables")
    images = [r for r in doc.part.rels.values() if "image" in r.reltype]
    check("DOCX: images embedded", len(images) >= len(figs), f"{len(images)} images for {len(figs)} figure captions")
    s = doc.sections[0]
    check("DOCX: A4 with 1-inch margins", abs(s.page_width.inches - 8.27) < 0.02 and abs(s.page_height.inches - 11.69) < 0.02
          and all(abs(m.inches - 1) < 0.01 for m in (s.left_margin, s.right_margin, s.top_margin, s.bottom_margin)))
    st = doc.styles["Normal"]
    check("DOCX: Normal style Times New Roman 12 pt, 1.5 spacing, justified",
          st.font.name == "Times New Roman" and st.font.size.pt == 12 and st.paragraph_format.line_spacing == 1.5
          and st.paragraph_format.alignment == 3)


def validate_pdf() -> None:
    if not PDF.exists():
        check("PDF: exported", False, "file missing")
        return
    d = fitz.open(PDF)
    check("PDF: exported", True, f"{d.page_count} pages")
    blank = [i + 1 for i, pg in enumerate(d) if not pg.get_text().strip().replace("Computer Networks Laboratory", "").strip(" 0123456789\n")
             and not pg.get_images()]
    check("PDF: no blank pages", not blank, f"blank pages: {blank}")
    full = "\n".join(pg.get_text() for pg in d)
    check("PDF: TOC populated with page numbers", bool(re.search(r"4 Testing and Results\s*\.*\s*\d+", full)))
    check("PDF: list of figures populated", "Figure 2.1: Architecture of the Distributed CDN Simulator" in full)
    check("PDF: no Word field errors", "Error!" not in full)
    check("PDF: newer than DOCX", PDF.stat().st_mtime >= DOCX.stat().st_mtime - 5)


def validate_pptx() -> None:
    prs = Presentation(str(PPTX))
    check("PPTX: exactly 8 slides", len(prs.slides) == 8, f"{len(prs.slides)} slides")
    text = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    check("PPTX: project title", TITLE in text)
    check("PPTX: faculty Dr. Swati", "Dr. Swati" in text)
    check("PPTX: Group 2", "Group 2" in text)
    for name, prn in TEAM:
        check(f"PPTX: {name} – {prn}", f"{name} – {prn}" in text)
    leftovers = [p for p in INSTRUCTION_PATTERNS if re.search(p, ALLOWED_PLACEHOLDERS.sub("", text))]
    check("PPTX: no template instructions left", not leftovers, ", ".join(leftovers))
    notes = "\n".join(s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide)
    check("PPTX: template note instructions replaced", "Replace all bracketed" not in notes)
    sw, shh = prs.slide_width, prs.slide_height
    outside = [(i + 1, sh.name) for i, s in enumerate(prs.slides) for sh in s.shapes
               if sh.left < 0 or sh.top < 0 or sh.left + sh.width > sw or sh.top + sh.height > shh]
    check("PPTX: all shapes inside the slide", not outside, str(outside))
    overlaps = []
    for i, s in enumerate(prs.slides):
        boxes = [(sh.name, sh.left, sh.top, sh.left + sh.width, sh.top + sh.height) for sh in s.shapes
                 if sh.shape_type == 13 and sh.name != "Picture 3"]  # inserted pictures (template logo excluded)
        texts = [(sh.name, sh.left, sh.top, sh.left + sh.width, sh.top + sh.height) for sh in s.shapes if sh.has_text_frame
                 and sh.text_frame.text.strip()]
        for b in boxes:
            for t in texts:
                if b[1] < t[3] and t[1] < b[3] and b[2] < t[4] and t[2] < b[4]:
                    overlaps.append((i + 1, b[0], t[0]))
    check("PPTX: inserted images do not overlap text boxes", not overlaps, str(overlaps))
    ref_box = next((sh for sh in prs.slides[-1].shapes if sh.has_text_frame and sh.text_frame.text.startswith("[1]")), None)
    refs = [ln for ln in ref_box.text_frame.text.split("\n") if re.match(r"^\[\d+\]", ln)] if ref_box else []
    check("PPTX: references on slide 8", len(refs) >= 5, f"{len(refs)} references")


def validate_templates() -> None:
    for path, digest in TEMPLATES.items():
        actual = hashlib.sha256(path.read_bytes()).hexdigest()
        check(f"Template unchanged: {path.name}", actual == digest, actual[:16])


def main() -> int:
    validate_templates()
    validate_docx()
    validate_pdf()
    validate_pptx()
    width = max(len(n) for n, _, _ in results)
    for name, ok, detail in results:
        print(f"{'PASS' if ok else 'FAIL'}  {name.ljust(width)}  {detail}")
    failed = sum(1 for _, ok, _ in results if not ok)
    print(f"\n{len(results) - failed}/{len(results)} checks passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
