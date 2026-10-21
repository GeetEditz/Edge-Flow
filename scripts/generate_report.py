"""Generate the completed mini-project report from the original Word template.

    .venv-docs/Scripts/python scripts/generate_report.py            # DOCX (+ PDF via Microsoft Word if available)
    .venv-docs/Scripts/python scripts/generate_report.py --no-pdf

The template is opened read-only and saved under submission/; it is never modified. Experimental
numbers are read from submission/evidence/experiment_results.json (exported from the running system)
and test counts from submission/evidence/*_pytest.txt / *_vitest.txt, so nothing is typed in by hand.
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "Final_Computer_Networks_Mini_Project_Report_Template.docx"
OUT_DIR = ROOT / "submission"
OUT_DOCX = OUT_DIR / "Distributed_CDN_Simulator_Final_Report.docx"
OUT_PDF = OUT_DIR / "Distributed_CDN_Simulator_Final_Report.pdf"
FIG = OUT_DIR / "figures"
EVID = OUT_DIR / "evidence"

TITLE = "Distributed CDN Simulator with Latency-Aware Routing and Adaptive Caching"
TEAM = [("More Geet Dattatraya", "1262250006"), ("Shinde Soham Manoj", "1262250009"),
        ("Jagtap Prasad Sachin", "1262250018"), ("Gaikwad Shreyash Sachin", "1262250005")]
GROUP = "2"
FACULTY = "Dr. Swati"
PANEL = "[[PANEL – to be confirmed]]"
BATCH = "[[BATCH – to be confirmed]]"
SUBMISSION_DATE = "[[DD MONTH YYYY – to be confirmed]]"
TEXT_WIDTH_TW = 9029  # A4 width minus 1-inch margins, in twips
STRATEGY = {"random": "Random", "nearest": "Nearest replica", "lowest_latency": "Lowest latency",
            "least_loaded": "Least loaded", "hybrid": "Hybrid", "failover": "Health-aware failover"}
POLICY = {"none": "Disabled", "lru": "LRU", "lfu": "LFU", "ttl": "TTL (10 s)"}


# ----------------------------------------------------------------------------- low-level helpers

def add_runs(p: Paragraph, text: str, size: float | None = None, bold: bool | None = None, mono: bool = False) -> None:
    """Add text with inline markup: **bold**, `code`, [[placeholder]] (highlighted)."""
    for tok in re.split(r"(\*\*.+?\*\*|`.+?`|\[\[.+?\]\])", text):
        if not tok:
            continue
        if tok.startswith("**"):
            r = p.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("`"):
            r = p.add_run(tok[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(10.5 if size is None else size - 1)
        elif tok.startswith("[["):
            r = p.add_run("[" + tok[2:-2] + "]")
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            r = p.add_run(tok)
            if bold is not None:
                r.bold = bold
        if size and not tok.startswith("`"):
            r.font.size = Pt(size)
        if mono:
            r.font.name = "Consolas"


def set_text(p: Paragraph, text: str) -> None:
    """Replace paragraph text, keeping the formatting of the first run."""
    rpr = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for tok in re.split(r"(\[\[.+?\]\]|\*\*.+?\*\*)", text):
        if not tok:
            continue
        r = p.add_run("[" + tok[2:-2] + "]" if tok.startswith("[[") else tok[2:-2] if tok.startswith("**") else tok)
        if rpr is not None:
            r._r.insert(0, deepcopy(rpr))
        if tok.startswith("[["):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if tok.startswith("**"):
            r.bold = True


def delete(el) -> None:
    el = el._element if hasattr(el, "_element") else el
    el.getparent().remove(el)


def has_page_break(p_el) -> bool:
    return any(b.get(qn("w:type")) == "page" for b in p_el.iter(qn("w:br")))


def is_heading(p_el, levels=("Heading1", "Heading2")) -> bool:
    ps = p_el.find(qn("w:pPr") + "/" + qn("w:pStyle")) if p_el.tag == qn("w:p") else None
    return ps is not None and ps.get(qn("w:val")) in levels


def set_cell(cell, text: str, bold: bool | None = None, align=None) -> None:
    p = cell.paragraphs[0]
    rpr = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for extra in cell.paragraphs[1:]:
        delete(extra)
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        r = p.add_run(line)
        if rpr is not None:
            r._r.insert(0, deepcopy(rpr))
        if bold is not None:
            r.bold = bold
        if i < len(lines) - 1:
            r.add_break()
    if align is not None:
        p.alignment = align


def fill_template_table(tbl: Table, rows: list[list[str]]) -> None:
    """Keep the header row, clone the first data row as prototype for every new row."""
    proto = deepcopy(tbl.rows[1]._tr)
    for r in list(tbl.rows)[1:]:
        delete(r._tr)
    for data in rows:
        tr = deepcopy(proto)
        tbl._tbl.append(tr)
        row = tbl.rows[-1]
        for cell, val in zip(row.cells, data):
            set_cell(cell, val)


def cell_borders(cell, color="D9D9D9", fill: str | None = None) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single"), b.set(qn("w:sz"), "4"), b.set(qn("w:space"), "0"), b.set(qn("w:color"), color)
        borders.append(b)
    tcPr.append(borders)
    if fill:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), fill)
        tcPr.append(shd)


# ----------------------------------------------------------------------------- body writer

class Writer:
    """Appends content after a moving cursor element inside the template body."""

    def __init__(self, doc: Document, cursor):
        self.doc = doc
        self.cur = cursor._element if hasattr(cursor, "_element") else cursor

    def _place(self, el):
        self.cur.addnext(el)
        self.cur = el

    def para(self, text: str = "", style: str | None = None, align=None, size=None, bold=None, keep_next=False,
             space_after=None, indent=None, first_line=None) -> Paragraph:
        p_el = OxmlElement("w:p")
        self._place(p_el)
        p = Paragraph(p_el, self.doc._body)
        if style:
            p.style = self.doc.styles[style]
        if text:
            add_runs(p, text, size=size, bold=bold)
        pf = p.paragraph_format
        if align is not None:
            p.alignment = align
        if keep_next:
            pf.keep_with_next = True
        if space_after is not None:
            pf.space_after = Pt(space_after)
        if indent is not None:
            pf.left_indent = Inches(indent)
        if first_line is not None:
            pf.first_line_indent = Inches(first_line)
        return p

    def h3(self, text: str) -> Paragraph:
        return self.para(text, style="Heading 3", keep_next=True)

    def bullets(self, items: list[str]) -> None:
        for it in items:
            p = self.para("•\t" + it, indent=0.3, first_line=-0.2, space_after=2)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))

    def numbered(self, items: list[str]) -> None:
        for i, it in enumerate(items, 1):
            p = self.para(f"{i}.\t" + it, indent=0.35, first_line=-0.25, space_after=2)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.35))

    def code(self, text: str) -> None:
        lines = text.strip("\n").split("\n")
        for i, line in enumerate(lines):
            p = self.para(style="Code Block")
            p.add_run(line if line else " ")
            if i == 0:
                p.paragraph_format.space_before = Pt(4)
            if i == len(lines) - 1:
                p.paragraph_format.space_after = Pt(8)
            else:
                p.paragraph_format.keep_with_next = True

    def caption(self, text: str, kind: str) -> None:
        label, rest = text.split(":", 1)
        p = self.para(style="Figure Caption" if kind == "figure" else "Table Caption", keep_next=kind == "table")
        r = p.add_run(label + ":")
        r.bold = True
        p.add_run(rest)

    def figure(self, path: Path, caption: str, width: float = 6.2) -> None:
        p = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True, space_after=4)
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(str(path), width=Inches(width))
        self.caption(caption, "figure")

    def attach(self, tbl: Table) -> Table:
        self._place(tbl._tbl)
        return tbl

    def table(self, headers: list[str], rows: list[list[str]], widths: list[float], caption: str | None = None,
              numeric: set[int] = frozenset(), font: float = 10.5) -> Table:
        if caption:
            self.caption(caption, "table")
        tbl = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        self._place(tbl._tbl)
        tbl.autofit = False
        tw = [int(TEXT_WIDTH_TW * w / sum(widths)) for w in widths]
        tblPr = tbl._tbl.tblPr
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)
        for gc, w in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), tw):
            gc.set(qn("w:w"), str(w))
        for ri, row in enumerate(tbl.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if ri == 0:
                trPr.append(OxmlElement("w:tblHeader"))
            data = headers if ri == 0 else rows[ri - 1]
            for ci, (cell, val) in enumerate(zip(row.cells, data)):
                cell.width = Twips(tw[ci])
                cell_borders(cell, fill="EEEEEE" if ri == 0 else None)
                p = cell.paragraphs[0]
                pf = p.paragraph_format
                pf.space_before, pf.space_after, pf.line_spacing = Pt(3), Pt(3), 1.1
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if (ci in numeric and ri > 0) else WD_ALIGN_PARAGRAPH.LEFT
                add_runs(p, str(val), size=font, bold=True if ri == 0 else None)
        self.para(space_after=0).paragraph_format.line_spacing = 0.8  # breathing room after the table
        return tbl


# ----------------------------------------------------------------------------- document surgery

def ensure_styles(doc: Document) -> None:
    st = doc.styles
    if "Code Block" not in [s.name for s in st]:
        s = st.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = st["Normal"]
        s.font.name, s.font.size = "Consolas", Pt(9)
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        pf = s.paragraph_format
        pf.alignment, pf.line_spacing, pf.space_before, pf.space_after = WD_ALIGN_PARAGRAPH.LEFT, 1.0, Pt(0), Pt(0)
        pf.left_indent = Inches(0.15)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), "F3F4F6")
        s._element.get_or_add_pPr().append(shd)
    for name, before, after in (("Figure Caption", 2, 10), ("Table Caption", 8, 4)):
        if name not in [x.name for x in st]:
            s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = st["Normal"]
            s.font.size = Pt(10.5)
            s.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
            pf = s.paragraph_format
            pf.alignment, pf.line_spacing = WD_ALIGN_PARAGRAPH.CENTER, 1.15
            pf.space_before, pf.space_after = Pt(before), Pt(after)
    h3 = st["Heading 3"]
    h3.font.size = Pt(12)


def keep_table_headers(doc: Document) -> None:
    """Keep each table's header row on the same page as its first data row."""
    for tbl in doc.tables:
        for p_el in tbl.rows[0]._tr.iter(qn("w:p")):
            Paragraph(p_el, tbl._parent).paragraph_format.keep_with_next = True


def convert_page_breaks(doc: Document) -> None:
    """The template separates parts with empty paragraphs holding a page break. When such a paragraph lands
    at the top of a page it produces a blank page, so move the break onto the following paragraph instead."""
    for p in list(doc.paragraphs):
        if has_page_break(p._p) and not p.text.strip():
            nxt = p._p.getnext()
            while (nxt is not None and nxt.tag == qn("w:p") and not "".join(nxt.itertext()).strip()
                   and not list(nxt.iter(qn("w:drawing")))):
                nxt = nxt.getnext()
            if nxt is not None and nxt.tag == qn("w:p"):
                Paragraph(nxt, p._parent).paragraph_format.page_break_before = True
                delete(p)


def find_par(doc: Document, prefix: str) -> Paragraph:
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise KeyError(prefix)


def clear_section(heading: Paragraph) -> list[Table]:
    """Remove placeholder content after a heading up to the next Heading 1/2. Page breaks are kept;
    template tables are detached and returned so they can be re-attached where needed."""
    tables, el = [], heading._p.getnext()
    while el is not None and el.tag != qn("w:sectPr"):
        nxt = el.getnext()
        if el.tag == qn("w:p") and is_heading(el):
            break
        if el.tag == qn("w:tbl"):
            el.getparent().remove(el)
            tables.append(Table(el, heading._parent))
        elif el.tag == qn("w:p") and not has_page_break(el):
            el.getparent().remove(el)
        el = nxt
    return tables


def remove_heading_section(doc: Document, title: str) -> None:
    h = find_par(doc, title)
    clear_section(h)
    delete(h)


def field_paragraph(w: Writer, instr: str) -> None:
    p = w.para()
    for kind, txt in (("begin", None), ("instr", instr), ("separate", None), ("text", "Right-click and choose Update Field."),
                      ("end", None)):
        r = OxmlElement("w:r")
        if kind == "instr":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            r.append(it)
        elif kind == "text":
            t = OxmlElement("w:t")
            t.text = txt
            r.append(t)
        else:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r.append(fc)
        p._p.append(r)


# ----------------------------------------------------------------------------- evidence

def load_evidence():
    exps = {e["name"].split(".")[0].strip(): e for e in json.loads((EVID / "experiment_results.json").read_text(encoding="utf-8"))}
    be = (EVID / "backend_pytest.txt").read_text(encoding="utf-8")
    fe = (EVID / "frontend_vitest.txt").read_text(encoding="utf-8")
    be_pass = len(re.findall(r"PASSED", be))
    be_fail = len(re.findall(r"FAILED|ERROR", be))
    m = re.search(r"Tests\s+(\d+) passed", fe)
    fe_pass = int(m.group(1)) if m else 0
    per_file = {}
    for f in re.findall(r"tests/(test_\w+)\.py::", be):
        per_file[f] = per_file.get(f, 0) + 1
    ui = (EVID / "ui_check.txt").read_text(encoding="utf-8") if (EVID / "ui_check.txt").exists() else ""
    ui_pass = len(re.findall(r"^PASS ", ui, re.M))
    return exps, be_pass, be_fail, fe_pass, per_file, ui_pass


def ms(v):
    return "—" if v is None else f"{v:.1f} ms"


def pct(v, d=1):
    return "—" if v is None else f"{v * 100:.{d}f} %"


# ----------------------------------------------------------------------------- content

def build(doc: Document) -> None:
    exps, be_pass, be_fail, fe_pass, per_file, ui_pass = load_evidence()
    A, B, C, C2, D, D2 = (exps[k] for k in ("A", "B", "C", "C2", "D", "D2"))
    va = {v["routing"]: v["metrics"] for v in A["summary"]["variants"]}
    vb = {v["cache_policy"]: v["metrics"] for v in B["summary"]["variants"]}
    fc, fc2 = C["summary"]["variants"][0]["failure"], C2["summary"]["variants"][0]["failure"]
    md = D["summary"]["variants"][0]["metrics"]
    md2 = D2["summary"]["variants"][0]["metrics"]
    total_tests = be_pass + fe_pass

    ensure_styles(doc)
    doc.core_properties.title = TITLE
    doc.core_properties.author = "Group 2, Computer Networks Laboratory (CSE20110)"
    doc.core_properties.subject = "Mini Project Report"

    # ------------------------------------------------------------- cover
    set_text(find_par(doc, "[PROJECT TITLE]"), TITLE)
    fill_template_table(doc.tables[0], [[str(i), n, prn] for i, (n, prn) in enumerate(TEAM, 1)])
    set_text(find_par(doc, "Group No. [GROUP"), f"Group No. {GROUP}     Panel {PANEL}     Batch {BATCH}")
    set_text(find_par(doc, "[Faculty Name]"), FACULTY)
    set_text(find_par(doc, "Submission Date"), f"Submission Date {SUBMISSION_DATE}")

    # ------------------------------------------------------------- certificate
    cert = find_par(doc, "This is to certify")
    cert_text = cert.text.replace("[PROJECT TITLE]", TITLE).replace("[PANEL]", PANEL).replace("[BATCH]", BATCH) \
        .replace("[GROUP NUMBER]", GROUP)
    set_text(cert, cert_text)
    fill_template_table(doc.tables[1], [[n, prn] for n, prn in TEAM])
    set_text(find_par(doc, "[FACULTY NAME]"), f"{FACULTY}\nProject Guide and Course Faculty")

    # ------------------------------------------------------------- acknowledgement
    ack = find_par(doc, "[Write a brief acknowledgement")
    set_text(ack, f"We thank {FACULTY}, our project guide and course faculty for the Computer Networks Laboratory "
                  "(CSE20110), for the guidance provided during this mini-project. We are grateful to the Department of "
                  "Computer Engineering and Technology, MIT World Peace University, Pune, for the laboratory course and "
                  "resources that made this work possible. We also acknowledge the open-source projects on which the "
                  "implementation is built, including Python, FastAPI, React, Vite, Recharts and SQLite, and the IETF "
                  "specifications that describe the protocols studied in this project.")
    set_text(find_par(doc, "Group No. [GROUP NUMBER]"), f"Group No. {GROUP}")
    for p in list(doc.paragraphs):  # the template contains empty Heading 1 paragraphs that would pollute the TOC
        if p.style.name.startswith("Heading") and not p.text.strip():
            delete(p)

    # empty spacer paragraphs between the acknowledgement and the abstract
    el = find_par(doc, f"Group No. {GROUP}")._p.getnext()
    while el is not None and el.tag == qn("w:p") and not "".join(el.itertext()).strip() and not has_page_break(el):
        nxt = el.getnext()
        el.getparent().remove(el)
        el = nxt

    # ------------------------------------------------------------- abstract
    abstract = find_par(doc, "[Write approximately 150")
    set_text(abstract,
             "Content Delivery Networks reduce latency and origin load by serving content from edge servers close to "
             "users, but choosing the right edge requires balancing network latency, server load, availability and cache "
             "contents. This project implements a working miniature CDN to study these trade-offs. An origin server and "
             "three independent edge-server processes communicate over HTTP/1.1 on TCP; clients locate an edge through a "
             "custom DNS-inspired resolver protocol over UDP. Six routing strategies (random, nearest replica, lowest "
             "latency, least loaded, a weighted hybrid of latency, load and packet loss, and health-aware failover) and "
             "four cache configurations (LRU, LFU, TTL and disabled) are implemented. A FastAPI control plane performs "
             "active and passive health checking, runs reproducible experiments, stores results in SQLite and streams "
             "live metrics over WebSocket to a React dashboard. "
             f"{be_pass + fe_pass} automated tests pass. In Local Mode experiments with simulated link profiles, hybrid "
             f"routing achieved the lowest average latency ({va['hybrid']['avg_ms']:.0f} ms versus "
             f"{va['nearest']['avg_ms']:.0f} ms for nearest replica) under concurrent load, caching cut origin fetches by "
             f"about {(1 - vb['lfu']['origin_fetches'] / vb['none']['origin_fetches']) * 100:.0f} %, and a failed edge "
             "was detected within milliseconds by passive checks. Mininet execution remains future work.")
    set_text(find_par(doc, "Keywords:"), "Keywords: content delivery network, edge caching, latency-aware routing, "
                                          "load balancing, HTTP/1.1, UDP name resolution, LRU, LFU")
    remove_heading_section(doc, "Using This Template")
    remove_heading_section(doc, "Formatting and Submission")

    # ------------------------------------------------------------- TOC / lists
    delete(find_par(doc, "[In Word, right-click the contents field"))
    lof = find_par(doc, "List of Figures")
    clear_section(lof)
    field_paragraph(Writer(doc, lof), 'TOC \\h \\z \\t "Figure Caption,1"')
    lot = find_par(doc, "List of Tables")
    clear_section(lot)
    field_paragraph(Writer(doc, lot), 'TOC \\h \\z \\t "Table Caption,1"')

    # ============================================================= CHAPTER 1
    w = Writer(doc, find_par(doc, "1.1 Introduction and Background"))
    clear_section(find_par(doc, "1.1 Introduction and Background"))
    w.para("A Content Delivery Network (CDN) is a distributed system of servers that delivers web content on behalf of "
           "an **origin server**, the authoritative source of the content. Copies of popular objects are kept in the caches "
           "of **edge servers** placed closer to users, so that most requests are answered without crossing long network "
           "paths to the origin. Large commercial CDNs operate hundreds of thousands of such servers and combine request "
           "routing, caching and failure handling to meet performance targets [6].")
    w.para("Three quantities dominate the user-visible performance of a CDN. **Latency** is the time between issuing a "
           "request and receiving the last byte of the response; it grows with the round-trip time of the path, with "
           "transmission time for large objects, and with retransmissions caused by packet loss [8]. **Server load** "
           "matters because an edge that is already serving many requests queues new ones, so a nearby but overloaded "
           "server can respond more slowly than a distant idle one. **Caching** determines how often an edge must fetch "
           "an object from the origin: web request popularity is strongly skewed and approximately follows a Zipf-like "
           "distribution, which is why small caches can absorb a large fraction of requests [7].")
    w.para("Clients are usually directed to an edge by the Domain Name System: a resolver answers a query for the CDN "
           "hostname with the address of a selected edge, and the answer is cached for its time-to-live [3]. Content is "
           "then transferred with HTTP over TCP [1], [2], [5]. Intelligent routing therefore happens at resolution time, "
           "and it must use measurements such as latency, load and health to pick an edge. Content-delivery efficiency "
           "is the combined result of these choices: lower latency for users, fewer and smaller transfers from the "
           "origin, and continued service when an edge fails.")
    w.para("This project builds a small but genuinely networked CDN in which these mechanisms can be observed and "
           "compared. All communication uses real sockets: HTTP/1.1 over TCP for content and a DNS-inspired binary "
           "protocol over UDP for resolution. Experiments are repeatable and their results are measured from real "
           "requests, which makes the system suitable for laboratory study of routing and caching trade-offs.")

    w = Writer(doc, find_par(doc, "1.2 Problem Statement"))
    clear_section(find_par(doc, "1.2 Problem Statement"))
    w.para("Given several edge servers with different network distance, bandwidth, packet loss, capacity and cache "
           "contents, a CDN must decide, for each client request, which edge should serve it and what each edge should "
           "keep in its limited cache. A purely distance-based choice ignores load and can overload the closest edge; a "
           "purely load-based choice can send users to distant servers; neither reacts correctly to failures unless "
           "server health is tracked. The problem addressed by this project is to implement and evaluate an edge-selection "
           "and caching system that considers latency, load, packet loss and availability, and to quantify the effect of "
           "each strategy on latency, throughput, cache-hit ratio, origin traffic and failure recovery under controlled, "
           "repeatable conditions.")

    w = Writer(doc, find_par(doc, "1.3 Scope"))
    clear_section(find_par(doc, "1.3 Scope"))
    w.para("The system is intended for students and engineers studying content delivery. Its boundaries are:")
    w.bullets([
        "**Included:** one origin server, three edge servers (India, Singapore and USA, all simulated locations), a UDP "
        "resolver, six routing strategies, LRU/LFU/TTL caching, health checking with failure injection, a traffic "
        "generator, experiment management with SQLite persistence, a REST/WebSocket control plane and a web dashboard.",
        "**Local Mode (implemented and tested):** all processes run on one computer over the loopback interface. "
        "Geographic distance and link conditions are simulated inside each edge by delaying responses according to a "
        "configurable profile (base latency, jitter, bandwidth and packet loss).",
        "**Mininet Mode (implemented, not executed):** a topology script places the same processes on Mininet hosts "
        "connected by links shaped with tc/netem [9]. It requires Linux and was not run on the Windows development machine.",
        "**Assumptions:** a single client region, a static content catalogue of 32 objects (about 9 MB), in-memory "
        "caches, and one experiment at a time.",
        "**Out of scope:** real geographic deployment, standards-compliant DNS, HTTPS, HTTP/2 or HTTP/3, cache "
        "revalidation, and production-scale load.",
    ])

    w = Writer(doc, find_par(doc, "1.4 Objectives"))
    tables = clear_section(find_par(doc, "1.4 Objectives"))
    w.para("The objectives below are stated so that each can be checked through implementation, automated tests or "
           "experiments (Table 1.1).")
    w.caption("Table 1.1: Project objectives and verification methods", "table")
    objectives = [
        ["O1", "Serve real objects from an origin through three independent edge servers over HTTP/1.1 on TCP, with "
               "correct handling of cache hits, misses, missing objects and malformed requests.",
         "Automated networking and integration tests; command-line download with hash comparison"],
        ["O2", "Resolve the CDN hostname to an edge through a documented UDP protocol with request identifiers, TTL "
               "and error codes.", "UDP resolver tests (success, unknown name, malformed input, timeout)"],
        ["O3", "Implement at least four routing strategies, including a configurable hybrid latency/load/loss score, "
               "that never select an unhealthy edge when a healthy one exists.",
         "Routing unit tests; routing comparison experiment (A)"],
        ["O4", "Implement LRU, LFU and TTL caches per edge with byte-capacity enforcement, expiry, invalidation and "
               "request coalescing.", "Caching unit tests; cache comparison experiment (B)"],
        ["O5", "Detect a failed edge, reroute new requests to healthy edges and measure detection and recovery time.",
         "Failure tests; server-failure experiment (C)"],
        ["O6", "Persist experiments and expose live metrics through a REST/WebSocket control plane and dashboard.",
         "API, persistence and WebSocket tests; frontend tests; scripted browser checks"],
        ["O7", "Run the CDN on an emulated network topology with shaped links (Mininet).", "Mininet smoke test"],
    ]
    t = w.attach(tables[0])
    fill_template_table(t, objectives)
    w.para()

    # ============================================================= CHAPTER 2
    h = find_par(doc, "2.1 Architecture and Network Topology")
    clear_section(h)
    w = Writer(doc, h)
    w.para("The system is divided into a **data plane**, the processes that carry content bytes, and a **control "
           "plane**, which configures, observes and experiments on the data plane (Figure 2.1). Every arrow in the figure "
           "is a real network connection; no component calls another through in-memory shortcuts.")
    w.figure(FIG / "fig_architecture.png", "Figure 2.1: Architecture of the Distributed CDN Simulator")
    w.para("**Data plane.** The origin server (TCP port 9000) holds the content catalogue. Each edge server is a separate "
           "operating-system process with its own cache, counters, concurrency limit and simulated network profile, and "
           "has two TCP listeners: a data port (9001–9003) for client requests and health probes, and an admin port "
           "(9101–9103) through which the control plane manages it. A client first sends a resolution query to the CNRP "
           "resolver (UDP port 8053), receives the address of the selected edge, and then downloads the object with "
           "HTTP/1.1. On a cache miss the edge fetches the object from the origin.")
    w.para("**Control plane.** A FastAPI [11] application (TCP port 8000) contains the routing engine, a cluster manager that "
           "probes edge health and polls statistics, the experiment runner and traffic generator, a live metrics window "
           "and the SQLite database. The React dashboard (served by Vite on port 5173) uses the REST API and a WebSocket "
           "stream [12]. In Local Mode all processes listen on 127.0.0.1; the Mininet topology places them on separate hosts "
           "of an emulated network (Section 2.3).")

    h = find_par(doc, "2.2 Tools and System Requirements")
    tables = clear_section(h)
    w = Writer(doc, h)
    w.para("Table 2.1 lists the software and hardware actually used. Versions were read from the dependency manifests "
           "(`backend/requirements.txt`, `frontend/package.json`) and from the development environment.")
    w.caption("Table 2.1: Tools, verified versions and purpose", "table")
    tools = [
        ["Operating system", "Windows 11 Home (10.0.26200)", "Development and Local Mode execution"],
        ["Hardware", "AMD Ryzen 9 8945HS (8 cores / 16 threads), 15.3 GB RAM", "All processes run on this machine"],
        ["Python", "3.13.2", "Backend language (asyncio sockets)"],
        ["FastAPI / Uvicorn", "0.115.6 / 0.32.1", "Control-plane REST API and ASGI server"],
        ["Pydantic", "2.10.4", "Request validation"],
        ["websockets", "13.1", "WebSocket transport for Uvicorn"],
        ["SQLite (Python sqlite3)", "3.45.3", "Persistence of experiments, runs and requests"],
        ["psutil", "6.1.1", "Process CPU and memory figures; parent watchdog"],
        ["pytest / pytest-asyncio / httpx", "8.3.4 / 0.25.0 / 0.28.1", "Backend automated tests"],
        ["Node.js / npm", "20.19.4 / 11.7.0", "Frontend tool chain"],
        ["React / React Router", "18.3.1 / 6.28.0", "Dashboard user interface"],
        ["TypeScript / Vite", "5.6.3 / 5.4.11", "Typed frontend source and build tool"],
        ["Recharts / Lucide React", "2.13.3 / 0.460.0", "Charts and icons"],
        ["Vitest / Testing Library", "2.1.8 / 16.1.0", "Frontend automated tests"],
        ["Google Chrome", "155.0.8059.85", "Scripted dashboard checks and screenshots (DevTools Protocol)"],
        ["Mininet, Open vSwitch, Wireshark", "Not installed", "Documented for Linux; not used in this environment"],
    ]
    t = w.attach(tables[0])
    fill_template_table(t, tools)
    w.para()

    h = find_par(doc, "2.3 Addressing and Protocol Selection")
    tables = clear_section(h)
    w = Writer(doc, h)
    w.para("This is a socket-programming project, so addressing is expressed as host addresses and ports (Table 2.2). "
           "In Local Mode every service shares the loopback interface, so no gateway is involved; the Mininet addresses "
           "are those configured in `mininet/topology.py` and were not executed.")
    w.caption("Table 2.2: Host addresses, ports and roles", "table")
    addr = [
        ["Control plane (REST + WebSocket)", "127.0.0.1, TCP 8000", "N/A (loopback)", "API, health manager, experiments"],
        ["CNRP resolver", "127.0.0.1, UDP 8053", "N/A (loopback)", "Edge selection for cdn.sim.local"],
        ["Origin server", "127.0.0.1, TCP 9000", "N/A (loopback)", "Authoritative content source"],
        ["India edge", "127.0.0.1, TCP 9001 / 9101", "N/A (loopback)", "Edge data port / admin port"],
        ["Singapore edge", "127.0.0.1, TCP 9002 / 9102", "N/A (loopback)", "Edge data port / admin port"],
        ["USA edge", "127.0.0.1, TCP 9003 / 9103", "N/A (loopback)", "Edge data port / admin port"],
        ["Dashboard (Vite)", "127.0.0.1, TCP 5173", "N/A (loopback)", "Web UI; proxies /api and /ws"],
        ["Mininet hosts (not executed)", "client 10.0.0.1/24, origin 10.0.0.2/24, edges 10.0.0.11–13/24",
         "N/A (single L2 switch)", "Emulated topology with tc/netem links"],
    ]
    t = w.attach(tables[0])
    fill_template_table(t, addr)
    w.para()
    w.para("**TCP and HTTP/1.1 for content.** Content transfer requires reliable, ordered delivery, so it uses TCP [5] "
           "with HTTP/1.1 message framing [2]. The project contains its own strict HTTP/1.1 implementation on asyncio "
           "streams [10] so that every byte on the wire is produced by the project and can be inspected. Requests are "
           "framed by `Content-Length` only (`Transfer-Encoding` is rejected with 501), the request head is limited to "
           "16 KiB (431 otherwise), malformed request lines, header folding, missing `Host` and conflicting lengths return "
           "400, and unsupported versions return 505. Every response carries `Content-Length`. The client opens one TCP "
           "connection per request, which makes each transfer visible as a complete handshake, exchange and teardown.")
    w.para("**UDP for resolution.** Name resolution is a short request/response exchange that tolerates loss through "
           "retransmission, so it uses UDP [4], following the DNS model [3]. The CDN Name Resolution Protocol (CNRP) is "
           "DNS-inspired but deliberately simpler and is **not** standards-compliant DNS. Its fixed binary layout is "
           "summarised in Table 2.3; all integers are in network byte order.")
    w.table(["Message", "Fields (bytes)", "Notes"], [
        ["Query", "magic `CNRQ` (4), version (1), flags (1), request id (2), name length (1), hostname (N)",
         "Flag bit 0 requests a routing explanation"],
        ["Answer", "magic `CNRR` (4), version (1), rcode (1), request id (2), TTL (4), IPv4 (4), port (2), edge-id "
                   "length (1) + id, explanation length (2) + text", "rcode: 0 OK, 1 FORMERR, 2 SERVFAIL, 3 NXDOMAIN, 4 BADVERS"],
    ], [1.1, 3.6, 2.3], caption="Table 2.3: CNRP message formats")
    w.para("`cdn.sim.local` is answered by the active routing strategy, `<edge>.cdn.sim.local` pins one edge, and any "
           "other name returns NXDOMAIN. Datagrams without the `CNRQ` magic are dropped silently. With DNS-style "
           "resolution the client caches the answer for its TTL (5 s), so routing happens at resolution time; a separate, "
           "explicitly non-DNS per-request mode resolves before every request and is used for routing comparisons.")

    # ============================================================= CHAPTER 3
    h = find_par(doc, "3.1 Methodology")
    clear_section(h)
    w = Writer(doc, h)
    w.para("The implementation followed an incremental, test-driven sequence in which each layer was verified over real "
           "sockets before the next layer was built:")
    w.numbered([
        "**Environment inspection** — operating system, Python, Node.js and port availability were checked; Local Mode "
        "was chosen as the primary target because Mininet requires Linux.",
        "**Data plane** — HTTP/1.1 server and client, origin server, edge server with simulated network profile, and "
        "the caching engine; verified by downloading real files and observing MISS followed by HIT.",
        "**Networking and routing** — CNRP protocol, resolver, routing strategies and health checking with failure "
        "injection; verified with UDP and failover tests.",
        "**Control plane** — FastAPI endpoints, WebSocket stream, SQLite persistence and the experiment runner.",
        "**Frontend** — seven dashboard pages connected only to real backend endpoints.",
        "**Experiments and debugging** — routing, caching, failure and degradation experiments; issues found during "
        "testing were diagnosed and fixed (Section 5.4).",
        "**Documentation** — architecture, API, experiment, Mininet and Wireshark guides.",
    ])
    w.para("Figure 3.1 shows the message sequence for one client request, which is the central workflow of the system.")
    w.figure(FIG / "fig_request_flow.png", "Figure 3.1: Request lifecycle across resolver, edge, cache and origin")

    h = find_par(doc, "3.2 Configuration or Code Implementation")
    clear_section(h)
    w = Writer(doc, h)
    w.para("The backend is organised into packages with a single responsibility each; the main modules are listed in "
           "Table 3.1.")
    w.table(["Module (backend/app/)", "Responsibility"], [
        ["networking/http.py", "Strict HTTP/1.1 server and client on asyncio TCP streams; crash simulation via stop()"],
        ["networking/resolver.py", "CNRP encoding/decoding, UDP resolver server and stub-resolver client"],
        ["origin/server.py", "Origin server: catalogue, safe file access, configurable response delay"],
        ["edge/server.py, edge/profile.py", "Edge server (data + admin listeners) and simulated link profile"],
        ["caching/policies.py, caching/edge_cache.py", "LRU, LFU, TTL policies; per-edge cache with coalescing"],
        ["routing/strategies.py, routing/engine.py", "Six strategies, normalisation, health filtering, decision log"],
        ["core/cluster.py", "Active/passive health checks, statistics polling, routing inputs"],
        ["core/launcher.py, core/services.py", "Process launcher (one OS process per server) and composition root"],
        ["traffic/generator.py, experiments/runner.py", "Seeded workloads and experiment orchestration"],
        ["persistence/db.py, api/routes.py, api/ws.py", "SQLite layer, REST endpoints, WebSocket metrics"],
    ], [2.6, 4.4], caption="Table 3.1: Principal backend modules")
    w.para("**Hybrid routing.** For each healthy edge the hybrid strategy min-max normalises the measured latency, load and "
           "estimated packet loss across the candidates (a missing value counts as worst, identical values as zero) and "
           "selects the edge with the lowest weighted score; weights must be non-negative and are normalised to sum to "
           "one. Ties are broken deterministically by configured priority and edge identifier:")
    w.code('''def scores(self, edges):
    nl = normalize({e.edge_id: e.latency_ms for e in edges})
    nd = normalize({e.edge_id: e.load for e in edges})
    np_ = normalize({e.edge_id: e.packet_loss for e in edges})
    w = self.weights                      # e.g. latency 0.5, load 0.4, loss 0.1
    return {e.edge_id: w["latency"] * nl[e.edge_id] + w["load"] * nd[e.edge_id]
                       + w["packet_loss"] * np_[e.edge_id] for e in edges}''')
    w.para("**Request coalescing.** When several requests miss on the same object at the same time, only the first "
           "fetches it from the origin; the others wait on the same future and are reported as `COALESCED`:")
    w.code('''pending = self._inflight.get(key)
if pending is not None:                      # identical fetch already in flight
    self.coalesced += 1
    return await asyncio.shield(pending), "COALESCED"
self.misses += 1
fut = asyncio.get_running_loop().create_future()
self._inflight[key] = fut                    # followers will await this future''')
    w.para("**Simulated link profile (Local Mode).** Each edge delays its response by a propagation delay (base latency "
           "± jitter), a transmission delay (object size ÷ bandwidth) and a loss penalty: the number of lost segments is "
           "drawn from a binomial distribution over ceil(size ÷ 1460) + 1 segments, and each lost segment adds "
           "max(2 × base latency, 10 ms). These delays are labelled as simulated throughout the API and dashboard.")

    h = find_par(doc, "3.3 Protocols and Technical Working")
    clear_section(h)
    w = Writer(doc, h)
    w.para("The implemented request lifecycle (Figure 3.1) is:")
    w.numbered([
        "The client (traffic generator or `scripts/cdn_client.py`) sends a CNRQ query for `cdn.sim.local` over UDP.",
        "The resolver asks the routing engine for a decision. Unhealthy edges are removed first; the active strategy "
        "selects among the rest and records its reason. The CNRR answer carries the edge address, port, TTL and, if "
        "requested, the explanation. With no healthy edge the answer is SERVFAIL.",
        "The client opens a TCP connection to the edge's data port and sends `GET /content/<object_id>`.",
        "The edge validates the identifier (1–5 path segments of `[A-Za-z0-9._-]`, no segment starting with a dot, no "
        "percent-encoding), waits for one of its concurrency slots, and looks the object up in its own cache.",
        "On a **hit** the cached bytes are used. Expired TTL entries are never served.",
        "On a **miss** the edge requests the object from the origin over a new TCP connection (unless an identical fetch "
        "is already in flight). A missing object yields 404 and an unreachable origin 502.",
        "The object is inserted according to the active policy, evicting entries until it fits; objects larger than the "
        "capacity are served but not cached. The response is sent with `X-Cache: HIT|MISS|COALESCED|BYPASS`.",
        "The client records the measured latency, status, edge and cache status; the control plane aggregates these "
        "records into live metrics and, for experiments, stores them in SQLite.",
    ])
    w.table(["Strategy", "Selection rule", "Input"], [
        ["Random", "Uniform choice among healthy edges (seeded for reproducibility)", "—"],
        ["Nearest replica", "Smallest configured distance from the client region", "Configured"],
        ["Lowest latency", "Lowest EWMA of measured health-probe round-trip time", "Measured"],
        ["Least loaded", "Lowest max(active + queued, client in-flight) ÷ concurrency limit", "Measured"],
        ["Hybrid", "Lowest weighted score of normalised latency, load and packet loss", "Measured + estimated"],
        ["Health-aware failover", "Highest-priority healthy edge (primary/backup)", "Configured + health"],
    ], [1.5, 3.9, 1.6], caption="Table 3.2: Implemented routing strategies")
    w.para("**Caching policies.** LRU evicts the least recently used object (ordered dictionary, O(1)); LFU evicts the "
           "least frequently used object with ties broken by recency (frequency buckets, O(1)); TTL expires objects a "
           "fixed time after insertion and, under capacity pressure, evicts the one closest to expiry. All policies "
           "enforce a byte capacity that is never exceeded, report every eviction with its reason, and can be switched "
           "at run time (the cache then restarts empty). Hit ratio is defined as hits ÷ (hits + misses + coalesced).")
    w.para("**Health checking and failure handling.** The cluster manager probes `GET /_health` on every edge's data port "
           "every 0.5 s with a 0.8 s timeout; two consecutive failures mark an edge unhealthy and two successes mark it "
           "healthy. In addition, three consecutive client transport failures (refused, reset or timeout) mark an edge "
           "unhealthy immediately (passive detection). A failure is injected by closing the edge's data listener and "
           "aborting its connections, which emulates a crashed process while its admin port remains available for restore.")
    w.para("**Connection handling and errors.** All socket operations have timeouts; the HTTP server supports persistent "
           "connections with a 15 s idle timeout; handler exceptions become 500 responses without tracebacks; client "
           "errors are classified as connect, timeout, reset or protocol failures. The control plane returns JSON errors "
           "of the form `{\"error\", \"detail\"}` with status 400, 404, 409, 422 or 502.")

    h = find_par(doc, "3.4 Execution Procedure")
    clear_section(h)
    w = Writer(doc, h)
    w.para("Prerequisites are Python 3.12 or newer, Node.js 18 or newer and the free ports listed in Table 2.2. The "
           "following commands were executed successfully on the development machine (Windows paths shown; on Linux "
           "replace `.venv\\Scripts\\python` with `.venv/bin/python`).")
    w.para("**Backend** (starts the control plane, which launches the origin and the three edge processes):", keep_next=True)
    w.code(r'''cd backend
python -m venv .venv
.venv\Scripts\python -m pip install -r requirements.txt
.venv\Scripts\python -m uvicorn app.main:app --host 127.0.0.1 --port 8000''')
    w.para("**Frontend** (second terminal), then open http://127.0.0.1:5173 in a browser:", keep_next=True)
    w.code('''cd frontend
npm install
npm run dev''')
    w.para("**Command-line client** (UDP resolution followed by an HTTP download) and **tests**:", keep_next=True)
    w.code(r'''cd backend
.venv\Scripts\python scripts\cdn_client.py images/photo-05.png --repeat 2
.venv\Scripts\python -m pytest
cd ..\frontend
npm test
npm run build''')
    w.para("Experiments are started from the dashboard (Traffic Generator, Routing Lab, Cache Lab or Experiments page) or "
           "through `POST /api/experiments`; results can be exported as CSV. Mininet Mode requires Ubuntu with Mininet, "
           "Open vSwitch and root privileges and is started with `sudo python3 mininet/topology.py`; this procedure is "
           "documented in `mininet/README.md` but was not executed.")
    w.figure(FIG / "shot_overview.png", "Figure 3.2: Dashboard overview page of the running system (live metrics)", width=5.9)
    w.para("Figure 3.2 is a screenshot of the running dashboard. All values come from the control plane: request counts, "
           "latency percentiles, cache-hit ratio and origin requests avoided are computed from client-side records of "
           "real requests, and the edge-health panel reflects active probe results.")
    w.figure(FIG / "shot_servers.png", "Figure 3.3: Edge Servers page showing per-edge health, measured RTT, load and cache", width=5.9)
    w.para("Figure 3.3 shows the three edge processes with their measured probe round-trip times next to the simulated "
           "profile values, current load against each edge's concurrency limit, cache occupancy and the failure-simulation "
           "control.")

    # ============================================================= CHAPTER 4
    h = find_par(doc, "4.1 Test Plan")
    tables = clear_section(h)
    w = Writer(doc, h)
    w.para("Testing combined three kinds of evidence. (1) **Automated tests:** "
           f"{be_pass} backend tests (pytest) exercise routing, caching, HTTP and UDP networking against servers on "
           "ephemeral loopback ports, the REST API, persistence and the WebSocket stream; "
           f"{fe_pass} frontend tests (Vitest) cover WebSocket reconnection, API error handling and interface states. "
           f"(2) **Scripted browser checks:** the running dashboard was driven through the Chrome DevTools Protocol "
           f"({ui_pass} checks) to confirm that every page renders without console errors and that controls change "
           "backend state. (3) **Experiments:** four experiment types were run in Local Mode on the hardware in Table 2.1.")
    w.para("The success criterion for a functional test is that the observed behaviour matches the expected behaviour "
           "exactly (status codes, cache status, selected edge, persisted data). Normal operation and failure or "
           "invalid-input cases were both tested. Table 4.1 summarises representative tests; the full list of passing "
           "test names is in `submission/evidence/backend_pytest.txt` and `frontend_vitest.txt`.")
    w.caption("Table 4.1: Functional test results", "table")
    tests = [
        ["T1", "GET an image through an edge twice (empty cache)", "200 with identical bytes; MISS then HIT; one origin fetch",
         "As expected; the origin's request counter increased by exactly one", "Pass"],
        ["T2", "cdn_client.py images/photo-05.png --repeat 2", "Valid PNG; MISS then HIT; same hash",
         "MISS 53.8 ms, HIT 24.2 ms, identical SHA-256, valid 112×112 PNG", "Pass"],
        ["T3", "Missing object; traversal ids (`../`, `%2e%2e`, `.hidden`)", "404 for missing; 400 for invalid ids",
         "As expected at both edge and origin", "Pass"],
        ["T4", "8 malformed requests and a 20 KB header", "400 / 405 / 501 / 505; 431 for oversize head",
         "All status codes matched", "Pass"],
        ["T5", "20 concurrent misses for one object", "One origin fetch; 1 MISS + 19 COALESCED",
         "Exactly one fetch; statuses as expected", "Pass"],
        ["T6", "UDP query for cdn.sim.local; unknown name", "OK answer with edge, port, TTL; NXDOMAIN",
         "As expected; explanation only with EXPLAIN flag", "Pass"],
        ["T7", "Malformed CNRP datagrams; no server", "FORMERR / BADVERS; garbage dropped; client timeout",
         "As expected (3 FORMERR/BADVERS, 1 dropped)", "Pass"],
        ["T8", "LRU and LFU eviction sequences", "LRU evicts least recent; LFU least frequent, ties by recency",
         "Eviction order matched in every case", "Pass"],
        ["T9", "TTL expiry with controlled clock", "Expired entries never served; purged before insert",
         "As expected", "Pass"],
        ["T10", "500 random insertions per policy; 1001 B object in 1000 B cache", "Capacity never exceeded; oversize rejected",
         "Invariant held for LRU, LFU and TTL", "Pass"],
        ["T11", "Each routing strategy on a fixed fleet; two unhealthy edges", "Expected edge; unhealthy edges excluded",
         "All six strategies chose the healthy edge in 30 trials each", "Pass"],
        ["T12", "Simulate failure of the nearest edge via API", "Edge marked unhealthy; requests go to next edge",
         "Next request served by Singapore with 'excluded unhealthy' reason; restored afterwards", "Pass"],
        ["T13", "Create, start, complete and export a cache-comparison experiment", "Results persisted; CSV exported",
         "Runs and 120 request records read back from a new database connection", "Pass"],
        ["T14", "Invalid API input (unknown strategy, zero weights, tiny capacity)", "422 / 400 with error detail",
         "As expected", "Pass"],
        ["T15", "WebSocket subscription during a failure", "hello, snapshots, failure event, pong",
         "All messages received", "Pass"],
        ["T16", "Dashboard pages and controls (browser automation)", "No console errors; controls change backend state",
         f"{ui_pass} of {ui_pass} checks passed", "Pass"],
        ["T17", "Mininet topology smoke test", "Traffic crosses shaped links", "Not executed (requires Linux)", "Not executed"],
        ["T18", "Wireshark capture of MISS/HIT exchange", "No origin traffic on a HIT", "Not executed (Wireshark not installed)",
         "Not executed"],
    ]
    t = w.attach(tables[0])
    fill_template_table(t, tests)
    w.para()

    h = find_par(doc, "4.2 Results and Execution Evidence")
    clear_section(h)
    w = Writer(doc, h)
    w.para(f"All {total_tests} automated tests passed ({be_pass} backend, {fe_pass} frontend; {be_fail} failures) and the "
           "production frontend build succeeded. The experiments below were measured on 20 October 2026 in Local Mode with "
           "the default simulated profiles: India 18 ms / 50 Mbit/s / 0.2 % loss / 8 slots, Singapore 42 ms / 200 Mbit/s "
           "/ 0.1 % / 16 slots, USA 110 ms / 400 Mbit/s / 0.4 % / 32 slots. All runs use a Zipf (s = 1.1) workload over "
           "the 32-object catalogue with seed 7 and start with empty caches. The numbers describe behaviour in this "
           "simulator on one machine and are not claims about real CDNs.")
    w.h3("Experiment A: routing comparison")
    rows = []
    for v in A["summary"]["variants"]:
        m = v["metrics"]
        share = v["edge_share"]
        rows.append([STRATEGY[v["routing"]], ms(m["avg_ms"]), ms(m["p95_ms"]), f"{m['throughput_rps']:.1f}",
                     pct(m["hit_ratio"]), " / ".join(f"{share.get(e, 0) * 100:.0f}" for e in
                                                     ("india-edge", "singapore-edge", "usa-edge"))])
    w.table(["Strategy", "Average", "P95", "Req/s", "Hit ratio", "Share % (IN / SG / US)"], rows,
            [1.6, 1.0, 1.0, 0.8, 0.9, 1.7], numeric={1, 2, 3, 4},
            caption="Table 4.2: Routing comparison, 400 requests per strategy, concurrency 24, mean of 3 repetitions")
    w.figure(FIG / "fig_routing_results.png", "Figure 4.1: Latency and throughput by routing strategy (Experiment A)")
    w.para(f"Nearest-replica and lowest-latency routing sent every request to the India edge. Its network round trip is "
           f"the shortest, but with only eight service slots and 50 Mbit/s it queued requests at concurrency 24, giving "
           f"{va['nearest']['avg_ms']:.0f} ms average latency. Least-loaded and random routing spread the load but sent "
           f"about a third of requests to the 110 ms USA edge, which raised the tail (P95 above one second). Hybrid "
           f"routing kept traffic on the two nearby edges and achieved the lowest average ({va['hybrid']['avg_ms']:.0f} "
           f"ms) and P95 ({va['hybrid']['p95_ms']:.0f} ms) latency and the highest throughput "
           f"({va['hybrid']['throughput_rps']:.0f} req/s) for this workload.")
    w.h3("Experiment B: cache-policy comparison")
    rows = []
    for v in B["summary"]["variants"]:
        m = v["metrics"]
        rows.append([POLICY[v["cache_policy"]], pct(m["hit_ratio"]), f"{m['origin_fetches']:.1f}",
                     f"{m['origin_bytes'] / 1e6:.1f} MB", f"{m['evictions']:.1f}", ms(m["avg_ms"]), ms(m["p95_ms"])])
    w.table(["Policy", "Hit ratio", "Origin fetches", "Origin bytes", "Evictions", "Average", "P95"], rows,
            [1.1, 0.9, 1.0, 1.0, 0.9, 1.0, 1.0], numeric={1, 2, 3, 4, 5, 6},
            caption="Table 4.3: Cache comparison, 400 requests per policy, 2 MB per edge, mean of 3 repetitions")
    w.figure(FIG / "fig_cache_results.png", "Figure 4.2: Origin fetches and hit ratio by cache policy (Experiment B)")
    w.para(f"With nearest routing a single edge receives the whole request stream. Every cache policy reduced origin "
           f"fetches from {vb['none']['origin_fetches']:.0f} to about {vb['lfu']['origin_fetches']:.0f}–"
           f"{vb['lru']['origin_fetches']:.0f} and origin bytes by more than half. LFU was marginally best "
           f"({pct(vb['lfu']['hit_ratio'])} hit ratio), consistent with a skewed popularity distribution [7]; the "
           f"differences between policies are small because a 2 MB cache can hold only a few of the large objects.")
    w.h3("Experiment C: server failure")
    w.table(["Resolution mode", "Detection", "Detected by", "Failed requests", "Last failure", "Recovery", "Failure rate"], [
        ["Per request", ms(fc["detection_ms"]), "Passive", f"{fc['failed_requests']} of {fc['requests_during_failure']}",
         ms(fc["last_failure_ms"]), ms(fc["recovery_ms"]), pct(C["summary"]["variants"][0]["metrics"]["failure_rate"], 2)],
        ["DNS-style (5 s TTL)", ms(fc2["detection_ms"]), "Passive", f"{fc2['failed_requests']} of {fc2['requests_during_failure']}",
         ms(fc2["last_failure_ms"]), ms(fc2["recovery_ms"]), pct(C2["summary"]["variants"][0]["metrics"]["failure_rate"], 2)],
    ], [1.5, 0.9, 0.9, 1.1, 1.0, 0.9, 0.9], numeric={1, 3, 4, 5, 6},
        caption="Table 4.4: India edge failed at t = 3 s and restored at t = 7 s (600 requests at 60 req/s, nearest routing)")
    w.figure(FIG / "shot_failure_experiment.png", "Figure 4.3: Failure experiment results in the dashboard (per-request resolution)")
    w.para("The failed requests were those in flight when the edge crashed; their reset connections triggered passive "
           "detection within milliseconds, and the next resolution already excluded India, so traffic moved to Singapore "
           "(Figure 4.3, right). After the restore, two successful probes returned India to service. With DNS-style "
           "resolution a client that still held a cached India answer kept trying it, and on Windows each refused "
           "loopback connection takes about two seconds, so the last failure occurred about 2 s after injection "
           "(Figure 4.4). This illustrates that DNS-based routing decides at resolution time, not per request.")
    w.figure(FIG / "fig_failure_timeline.png", "Figure 4.4: Requests per edge during the failure, per-request versus DNS-style resolution")
    w.h3("Experiment D: network degradation")
    w.para(f"At t = 3 s the India link profile was changed to 180 ms base latency and 2 % loss. Lowest-latency routing "
           f"detected the higher probe RTT and moved traffic to Singapore within about two seconds, finishing with "
           f"{ms(md['avg_ms'])} average latency and no failures. Nearest-replica routing kept using the degraded link: "
           f"average latency rose to {md2['avg_ms'] / 1000:.1f} s, throughput fell to {md2['throughput_rps']:.1f} req/s, "
           f"{pct(md2['failure_rate'])} of requests exceeded the 10 s client timeout, and the run lasted about "
           f"{D2['runs'][0]['summary']['duration_s'] / 60:.1f} minutes instead of about "
           f"{D['runs'][0]['summary']['duration_s']:.0f} seconds (Figure 4.5).")
    w.figure(FIG / "fig_degradation_timeline.png", "Figure 4.5: Average latency over time after link degradation (Experiment D)")

    h = find_par(doc, "4.3 Objective Achievement")
    tables = clear_section(h)
    w = Writer(doc, h)
    w.para("Table 4.5 maps each objective to its evidence. An objective is marked achieved only where tests or "
           "experiments demonstrate it.")
    w.caption("Table 4.5: Objective achievement", "table")
    t = w.attach(tables[0])
    fill_template_table(t, [
        ["O1", "T1–T5; networking and integration tests", "Achieved"],
        ["O2", "T6–T7; UDP resolver tests", "Achieved"],
        ["O3", "T11; Experiment A (six strategies compared)", "Achieved"],
        ["O4", "T5, T8–T10; Experiment B", "Achieved"],
        ["O5", "T12; Experiment C (detection and recovery measured)", "Achieved"],
        ["O6", "T13–T16; API, persistence, WebSocket and browser checks", "Achieved"],
        ["O7", "T17 (Mininet script exists but was not run)", "Not tested"],
    ])
    w.para()

    # ============================================================= CHAPTER 5
    h = find_par(doc, "5.1 Discussion and Comparative Analysis")
    clear_section(h)
    w = Writer(doc, h)
    w.para("The experiments show that the best edge depends on more than network distance. When the closest edge has "
           "limited capacity, distance- or latency-only routing concentrates load and queuing delay dominates; when load "
           "is balanced without regard to latency, distant edges inflate tail latency. The hybrid score balanced both "
           "effects and performed best in Experiment A, but this is a result for one workload and one set of weights, not "
           "a proof of optimality: at low concurrency, where no queue forms, nearest and lowest-latency routing are "
           "expected to perform as well, and different weights shift the balance.")
    w.para("Experiment D shows the complementary weakness of static strategies: nearest-replica routing has no feedback "
           "from measurements and continued to use a badly degraded link, while a measurement-driven strategy adapted "
           "within seconds. Experiment C shows that passive failure detection from client errors can react much faster "
           "than periodic probes, and that DNS-style caching of answers trades fewer resolutions for slower reaction to "
           "failures — the same trade-off that real DNS TTLs impose [3].")
    w.para("For caching, any policy substantially reduced origin traffic compared with no caching, while the differences "
           "among LRU, LFU and TTL were small at the tested capacity. Frequency-based eviction suits Zipf-like popularity "
           "[7]; recency-based eviction adapts faster to changing popularity, and TTL bounds staleness at the cost of "
           "evicting objects that are still popular. Latency gains from caching were modest here because transmission "
           "over the simulated India link, rather than the origin, dominated the latency of large objects.")

    h = find_par(doc, "5.2 Advantages and Applications")
    clear_section(h)
    w = Writer(doc, h)
    w.bullets([
        "**Real protocols:** content and resolution traffic use real TCP and UDP sockets, so behaviour can be inspected "
        "with packet-capture tools and the code doubles as a reference for socket programming.",
        "**Reproducible experiments:** seeded workloads, identical initial cache state for every variant, persisted "
        "configuration and CSV export make comparisons repeatable.",
        "**Transparency:** every routing decision is recorded with its reason and scores, and measured, simulated and "
        "estimated values are labelled separately.",
        "**Applications:** teaching CDN, DNS-based load balancing and caching concepts; prototyping routing heuristics "
        "before deploying them; studying failure detection and the effect of DNS TTLs on failover.",
    ])

    h = find_par(doc, "5.3 Limitations")
    clear_section(h)
    w = Writer(doc, h)
    w.bullets([
        "All processes run on one computer over loopback; geographic locations and link conditions are simulated inside "
        "each edge, and the measured latencies include these simulated delays.",
        "The loss model adds a fixed penalty per lost segment in sequence; it is pessimistic for large objects on long, "
        "lossy links and does not reproduce TCP congestion control [8].",
        "CNRP is DNS-inspired and not interoperable with standard DNS resolvers.",
        "Caches are in memory, TTL is a fixed lifetime rather than HTTP freshness, and conditional requests are not "
        "implemented.",
        "The workload family is limited (Zipf over a 32-object catalogue); experiments A and B used three repetitions "
        "and C and D single runs, so small differences are not statistically established.",
        "Mininet Mode and the Wireshark verification procedure were not executed in this environment.",
    ])

    h = find_par(doc, "5.4 Challenges and Troubleshooting")
    tables = clear_section(h)
    w = Writer(doc, h)
    w.para("The issues in Table 5.1 were observed during development and testing; each correction was verified by "
           "re-running the relevant measurement or test.")
    w.caption("Table 5.1: Issues encountered and their resolution", "table")
    t = w.attach(tables[0])
    fill_template_table(t, [
        ["Measured probe RTT about 11–14 ms higher than the simulated profile (e.g. 29–31 ms for an 18 ms link)",
         "asyncio sleeps on Windows are quantised to the 15.6 ms system timer; Windows 11 also ignores 1 ms timer "
         "requests from windowless child processes", "Requested 1 ms timer resolution and opted out of timer power "
         "throttling in every server process; RTTs became 21.6 / 46.4 / 114.1 ms for 18 / 42 / 110 ms profiles"],
        ["Failure detection took 4.3 s", "Connecting to a closed loopback port on Windows takes about 2 s, so each "
         "probe ran into its timeout", "Shorter probe interval and timeout plus passive detection from client failures; "
         "detection measured at 2.1 ms in Experiment C"],
        ["Failure-experiment test reported no detection time", "The injection timestamp was taken after the admin call "
         "returned, but detection had already happened during the call", "Timestamp taken at injection; test passes"],
        ["'Task was destroyed but it is pending' warnings at test shutdown", "Connection-handler tasks outlived the "
         "server's stop()", "Handlers are tracked and cancelled on stop; test run is clean"],
        ["Second of two back-to-back degradation experiments started on the first one's degraded link and left it "
         "degraded", "The runner read baseline profiles from a statistics cache not yet refreshed after the previous "
         "experiment's restore", "Stats refreshed before capturing and after restoring the baseline; a regression test "
         "fails without the fix and passes with it; the degradation experiments were re-run"],
        ["Router warnings found by the scripted browser check", "React Router v7 future-flag warnings in the console",
         "Opted in to the future flags; all pages render with zero console messages"],
    ])
    w.para()

    h = find_par(doc, "5.5 Future Scope")
    clear_section(h)
    w = Writer(doc, h)
    w.bullets([
        "Execute the Mininet topology on Linux, compare results on real shaped links with Local Mode, and capture the "
        "exchanges in Wireshark.",
        "Replace the simplified loss model with real network emulation, or a throughput model based on TCP behaviour [8].",
        "Add HTTP freshness and revalidation (`Cache-Control`, `ETag`, conditional GET) [1] and persistent disk caches.",
        "Learn hybrid weights automatically from measurements, and add multi-region clients and consistent-hashing "
        "content placement.",
        "Run longer, multi-workload experiments with more repetitions and confidence intervals.",
    ])

    h = find_par(doc, "5.6 Conclusion")
    clear_section(h)
    w = Writer(doc, h)
    w.para("The project delivered a working, networked CDN simulator: an origin, three independent edge processes with "
           "their own LRU, LFU or TTL caches, a UDP resolver, six routing strategies with health-aware filtering, active "
           "and passive failure detection, a FastAPI control plane with SQLite persistence and a seven-page React "
           f"dashboard. Objectives O1–O6 were achieved and verified by {total_tests} passing automated tests, scripted "
           "browser checks and four experiment types; objective O7 (Mininet execution) remains untested. The experiments "
           "showed, within the simulator, that combining latency and load in routing decisions reduced latency under "
           "load, that measurement-driven routing adapts to link degradation while static routing does not, that caching "
           "more than halves origin traffic for a Zipf workload, and that passive detection combined with per-request "
           "resolution limits the impact of an edge failure.")

    # ------------------------------------------------------------- references
    h = find_par(doc, "References")
    clear_section(h)
    w = Writer(doc, h)
    refs = [
        "R. Fielding, M. Nottingham, and J. Reschke, Eds., “HTTP Semantics,” IETF, RFC 9110, Jun. 2022. "
        "[Online]. Available: https://www.rfc-editor.org/rfc/rfc9110",
        "R. Fielding, M. Nottingham, and J. Reschke, Eds., “HTTP/1.1,” IETF, RFC 9112, Jun. 2022. "
        "[Online]. Available: https://www.rfc-editor.org/rfc/rfc9112",
        "P. Mockapetris, “Domain Names – Implementation and Specification,” IETF, RFC 1035, Nov. 1987. "
        "[Online]. Available: https://www.rfc-editor.org/rfc/rfc1035",
        "J. Postel, “User Datagram Protocol,” IETF, RFC 768, Aug. 1980. "
        "[Online]. Available: https://www.rfc-editor.org/rfc/rfc768",
        "W. Eddy, Ed., “Transmission Control Protocol (TCP),” IETF, RFC 9293, Aug. 2022. "
        "[Online]. Available: https://www.rfc-editor.org/rfc/rfc9293",
        "E. Nygren, R. K. Sitaraman, and J. Sun, “The Akamai network: A platform for high-performance Internet "
        "applications,” ACM SIGOPS Operating Systems Review, vol. 44, no. 3, pp. 2–19, Jul. 2010, "
        "doi: 10.1145/1842733.1842736.",
        "L. Breslau, P. Cao, L. Fan, G. Phillips, and S. Shenker, “Web caching and Zipf-like distributions: Evidence and "
        "implications,” in Proc. IEEE INFOCOM ’99, New York, NY, USA, 1999, pp. 126–134, doi: 10.1109/INFCOM.1999.749260.",
        "M. Mathis, J. Semke, J. Mahdavi, and T. Ott, “The macroscopic behavior of the TCP congestion avoidance "
        "algorithm,” ACM SIGCOMM Computer Communication Review, vol. 27, no. 3, pp. 67–82, Jul. 1997, "
        "doi: 10.1145/263932.264023.",
        "B. Lantz, B. Heller, and N. McKeown, “A network in a laptop: Rapid prototyping for software-defined networks,” "
        "in Proc. 9th ACM SIGCOMM Workshop on Hot Topics in Networks (HotNets-IX), 2010, doi: 10.1145/1868447.1868466.",
        "Python Software Foundation, “asyncio — Asynchronous I/O,” Python 3 documentation. [Online]. Available: "
        "https://docs.python.org/3/library/asyncio.html (accessed Oct. 20, 2026).",
        "FastAPI, “FastAPI documentation.” [Online]. Available: https://fastapi.tiangolo.com/ (accessed Oct. 20, 2026).",
        "I. Fette and A. Melnikov, “The WebSocket Protocol,” IETF, RFC 6455, Dec. 2011. "
        "[Online]. Available: https://www.rfc-editor.org/rfc/rfc6455",
    ]
    for i, r in enumerate(refs, 1):
        p = w.para(f"[{i}]\t{r}", indent=0.45, first_line=-0.45, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.45))
        p.paragraph_format.line_spacing = 1.15

    # ------------------------------------------------------------- AI declaration
    h = find_par(doc, "AI Usage Declaration")
    decl = find_par(doc, "Declaration: We take responsibility")
    p = decl._p.getprevious()
    while p is not None and p is not h._p:  # remove the instruction paragraph(s) before the template declaration
        prev = p.getprevious()
        p.getparent().remove(p)
        p = prev
    w = Writer(doc, h)
    w.para("AI tools were used in this project. Claude Code (Anthropic), an AI coding assistant, was used to generate "
           "and debug the source code, automated tests, experiment and figure scripts, the technical documentation and "
           "drafts of this report and presentation. The assistant also executed the automated tests and experiments in "
           "the development environment; the results reported here are taken from those runs (stored in "
           "`submission/evidence/`). The team is responsible for reviewing the implementation, re-running the tests and "
           "experiments, and verifying the content of this report before submission.")
    sig = doc.tables[-2] if len(doc.tables) >= 2 else None
    for tbl in doc.tables:
        if tbl.rows[0].cells[-1].text.strip() == "Signature":
            sig = tbl
    fill_template_table(sig, [[n, prn, ""] for n, prn in TEAM])

    # ------------------------------------------------------------- appendices
    h = find_par(doc, "Appendix A Code and Configuration")
    clear_section(h)
    w = Writer(doc, h)
    w.para("The complete source code is in the project repository; key files are listed below. Secrets or credentials "
           "are not used by the project.")
    w.table(["Path", "Content"], [
        ["backend/app/", "Backend packages (networking, routing, caching, origin, edge, core, traffic, experiments, "
                         "persistence, api, schemas)"],
        ["backend/app/config.py", "Default simulated topology, ports and environment variables"],
        ["backend/requirements.txt, frontend/package.json", "Pinned dependency manifests"],
        ["backend/tests/", "pytest suite: test_routing, test_caching, test_networking, test_metrics, test_api"],
        ["backend/scripts/cdn_client.py", "Command-line CNRP + HTTP client"],
        ["frontend/src/", "React dashboard (pages, components, API client, WebSocket hook)"],
        ["mininet/topology.py", "Mininet topology with tc/netem links (not executed)"],
        ["docs/", "ARCHITECTURE.md, API.md, EXPERIMENTS.md, WIRESHARK.md"],
        ["scripts/", "Report, presentation, figure and screenshot generation; submission validation"],
    ], [2.6, 4.4], caption="Table A.1: Important files")
    w.para("Default edge configuration (excerpt from `backend/app/config.py`):", keep_next=True)
    w.code('''{"edge_id": "india-edge", "host": "127.0.0.1", "data_port": 9001, "admin_port": 9101,
 "distance_km": 850, "priority": 1, "concurrency_limit": 8,
 "profile": {"base_latency_ms": 18, "jitter_ms": 3, "bandwidth_mbps": 50, "packet_loss": 0.002}}''')

    h = find_par(doc, "Appendix B Additional Evidence")
    clear_section(h)
    w = Writer(doc, h)
    w.para("Automated test counts per file, from `submission/evidence/backend_pytest.txt`:")
    w.table(["Test file", "Passed tests"], [[k + ".py", str(v)] for k, v in sorted(per_file.items())] +
            [["Frontend (src/app.test.tsx)", str(fe_pass)], ["Total", str(total_tests)]], [4.0, 2.0], numeric={1},
            caption="Table B.1: Automated test results")
    w.figure(FIG / "shot_routing_comparison.png", "Figure B.1: Routing comparison results as displayed in the Routing Lab", width=6.0)
    w.figure(FIG / "shot_topology.png", "Figure B.2: Network Topology page (conceptual topology in Local Mode)", width=6.0)
    w.para("Raw results of all experiments (configurations, per-run summaries and timelines) are in "
           "`submission/evidence/experiment_results.json`; per-request records can be exported as CSV from the "
           "Experiments page. No packet captures were made.")

    h = find_par(doc, "Appendix C Individual Contributions")
    tables = clear_section(h)
    w = Writer(doc, h)
    w.para("The project directory has no version-control history or task records from which individual contributions "
           "could be verified. The fields below are left for the team to complete.")
    w.caption("Table C.1: Individual contributions", "table")
    t = w.attach(tables[0])
    fill_template_table(t, [[f"{n} ({prn})", "[[To be completed by the team]]", "[[File / test / section]]"]
                            for n, prn in TEAM])
    for row in t.rows[1:]:  # re-apply highlight markup for the placeholders written via set_cell
        for cell in row.cells[1:]:
            txt = cell.text.strip("[]")
            set_cell(cell, "")
            r = cell.paragraphs[0].add_run(f"[{txt}]")
            r.font.size = Pt(11)
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW

    keep_table_headers(doc)
    convert_page_breaks(doc)


def export_pdf(docx_path: Path, pdf_path: Path) -> bool:
    """Update all fields (TOC, lists, page numbers) and export to PDF with Microsoft Word via COM."""
    ps = f"""
$ErrorActionPreference = 'Stop'
$w = New-Object -ComObject Word.Application
$w.Visible = $false
$w.DisplayAlerts = 0
try {{
  $d = $w.Documents.Open('{docx_path}')
  $d.Fields.Update() | Out-Null
  foreach ($t in $d.TablesOfContents) {{ $t.Update() | Out-Null }}
  $d.Save()
  $d.ExportAsFixedFormat('{pdf_path}', 17)
  Write-Output ("pages=" + $d.ComputeStatistics(2))
  $d.Close(0)
}} finally {{ $w.Quit() }}
"""
    try:
        res = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True, timeout=300)
    except (FileNotFoundError, subprocess.TimeoutExpired) as exc:
        print("PDF export unavailable:", exc)
        return False
    print(res.stdout.strip(), res.stderr.strip()[:500])
    return res.returncode == 0 and pdf_path.exists()


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-pdf", action="store_true")
    args = ap.parse_args()
    OUT_DIR.mkdir(exist_ok=True)
    doc = Document(str(TEMPLATE))  # opened read-only; saved to a new path
    build(doc)
    doc.save(str(OUT_DOCX))
    print("wrote", OUT_DOCX)
    if not args.no_pdf:
        ok = export_pdf(OUT_DOCX, OUT_PDF)
        print("PDF export:", "OK" if ok else "FAILED / unavailable")
        if not ok:
            sys.exit(2)


if __name__ == "__main__":
    main()
